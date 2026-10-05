# exportacao.py
import os
import io
import threading
import pandas as pd
from datetime import datetime
from tkinter import filedialog
import ttkbootstrap as tb
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
import sys
import database as db

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Image as RLImage

from componentes import ModernMessageBox

def executar_com_carregamento(app, funcao_pesada, titulo="Processando", mensagem="Aguarde..."):
    """Cria uma tela de carregamento segura que roda operações pesadas em background."""
    janela_load = tb.Toplevel(app)
    janela_load.title(titulo)
    janela_load.geometry("350x150")
    janela_load.resizable(False, False)
    janela_load.grab_set()
    
    # Centraliza a tela de carregamento
    janela_load.update_idletasks()
    x = app.winfo_x() + (app.winfo_width() // 2) - 175
    y = app.winfo_y() + (app.winfo_height() // 2) - 75
    janela_load.geometry(f"+{x}+{y}")

    tb.Label(janela_load, text=mensagem, font=("Helvetica", 11, "bold")).pack(pady=(25, 10))
    progresso = tb.Progressbar(janela_load, mode='indeterminate', bootstyle="success")
    progresso.pack(fill="x", padx=30)
    progresso.start(15)

    def tarefa_no_fundo():
        try:
            funcao_pesada()
            app.after(0, finalizar, True, None)
        except Exception as e:
            app.after(0, finalizar, False, str(e))

    def finalizar(sucesso, erro):
        progresso.stop()
        janela_load.destroy()
        if not sucesso:
            ModernMessageBox.showerror("Erro", f"Ocorreu um problema:\n{erro}", parent=app)

    # Inicia a thread secundária para não travar a tela principal
    threading.Thread(target=tarefa_no_fundo, daemon=True).start()

def gerar_excel_tabela(app, path, df_totais, df_resumo_raw, df_registros_export):
    """Lógica pesada para gerar o Excel da Aba Registros"""
    def trabalho_pesado():
        with pd.ExcelWriter(path, engine='openpyxl') as writer:
            # 1. Aba de Resumo Financeiro
            if df_totais is not None and not df_totais.empty:
                df_totais.to_excel(writer, sheet_name="Resumo Financeiro", index=False)
                ws = writer.sheets['Resumo Financeiro']
                ws.column_dimensions['A'].width = 25
                ws.column_dimensions['B'].width = 18
                ws.column_dimensions['C'].width = 18
                ws.column_dimensions['D'].width = 15

            # 2. Aba de Resumo Mensal
            if not df_resumo_raw.empty: 
                df_resumo_raw.to_excel(writer, sheet_name="Agrupado Mensal", index=False)
                
            # 3. Aba de Base de Dados Bruta
            if not df_registros_export.empty: 
                df_registros_export.to_excel(writer, sheet_name="Base de Dados", index=False)

        app.after(0, lambda: ModernMessageBox.showinfo("Sucesso", "Exportação concluída com sucesso!", parent=app))
    
    executar_com_carregamento(app, trabalho_pesado, "Exportando Excel", "Gerando planilhas, aguarde...")

def gerar_excel_dashboard(app, path, df_totais, df_detalhado):
    """Lógica pesada para gerar o Excel do Dashboard Financeiro"""
    def trabalho_pesado():
        with pd.ExcelWriter(path, engine='openpyxl') as writer:
            if df_totais is not None and not df_totais.empty:
                df_totais.to_excel(writer, sheet_name="Resumo Financeiro", index=False)
                ws = writer.sheets['Resumo Financeiro']
                fmt_moeda = 'R$ #,##0.00'
                
                # Ajusta a largura das colunas e formata como moeda
                for col in ws.columns:
                    col_name = col[0].value
                    length = max((len(str(cell.value)) for cell in col), default=10)
                    ws.column_dimensions[col[0].column_letter].width = length + 4
                    for cell in col[1:]:
                        if col_name in ["Crédito (Venda)", "Débito (Custo)"]: 
                            cell.number_format = fmt_moeda

            cols_export = ['Data', 'Parceiro', 'Resíduo', 'Estado', 'Destinação', 'Tipo', 'Peso (Kg)', 'Valor Unit.', 'Valor Total', 'Transporte']
            cols_finais = [c for c in cols_export if c in df_detalhado.columns]
            df_detalhado[cols_finais].to_excel(writer, sheet_name="Itens do Período", index=False)
            
        app.after(0, lambda: ModernMessageBox.showinfo("Sucesso", "Relatório Financeiro exportado!", parent=app))

    executar_com_carregamento(app, trabalho_pesado, "Exportando Dashboard", "Gerando arquivo Excel...")

def gerar_pdf_dashboard(app, path, dt_inicio, dt_fim, df_totais, img_estado, img_lucro, img_pizza):
    """Lógica pesada para desenhar e salvar o PDF"""
    def trabalho_pesado():
        row_total = df_totais[df_totais['Categoria'] == 'TOTAL GERAL'].iloc[0]
        v_vendas = row_total['Crédito (Venda)']
        v_custos = row_total['Débito (Custo)']
        v_lucro = v_vendas - v_custos
        v_peso = row_total['Peso (kg)']

        def fmt_moeda(v): return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        def fmt_kg(v): return f"{v:,.3f}".replace(",", "X").replace(".", ",").replace("X", ".")

        COR_PRIMARIA = colors.HexColor("#2C3E50")
        COR_SECUNDARIA = colors.HexColor("#18BC9C")
        COR_CINZA = colors.HexColor("#ECF0F1")

        def draw_header_footer(c, doc):
            c.saveState()
            w, h = landscape(A4)
            c.setFillColor(COR_PRIMARIA)
            c.rect(0, h - 2.5*cm, w, 2.5*cm, fill=1, stroke=0)
            c.setFillColor(colors.white)
            c.setFont("Helvetica-Bold", 18)
            c.drawString(2*cm, h - 1.5*cm, "RELATÓRIO DE GESTÃO DE RESÍDUOS")
            c.setFont("Helvetica", 10)
            c.drawString(2*cm, h - 2.1*cm, f"Período: {dt_inicio:%d/%m/%Y} a {dt_fim:%d/%m/%Y}")
            c.restoreState()

        doc = SimpleDocTemplate(path, pagesize=landscape(A4), topMargin=3.5*cm, bottomMargin=2.5*cm, leftMargin=1.5*cm, rightMargin=1.5*cm)
        elements = []
        styles = getSampleStyleSheet()
        h2 = ParagraphStyle('H2', parent=styles['Heading2'], textColor=COR_PRIMARIA, spaceBefore=15, spaceAfter=10)

        elements.append(Paragraph("Resumo de Indicadores", h2))
        kpi_data = [["RECEITA", "CUSTOS", "LUCRO", "VOLUME"],
                    [fmt_moeda(v_vendas), fmt_moeda(v_custos), fmt_moeda(v_lucro), fmt_kg(v_peso)]]
        cw = (landscape(A4)[0] - 4*cm) / 4
        kpi_table = Table(kpi_data, colWidths=[cw, cw, cw, cw])
        kpi_table.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'), 
            ('FONTNAME', (0,1), (-1,1), 'Helvetica-Bold'), 
            ('FONTSIZE', (0,1), (-1,1), 14)
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 0.5*cm))
        
        # Adiciona as imagens geradas dos gráficos do Matplotlib
        if img_estado: 
            elements.append(Table([[img_estado]], colWidths=[22*cm], style=[('ALIGN', (0,0), (-1,-1), 'CENTER')]))
        
        row2 = [img for img in [img_lucro, img_pizza] if img]
        if row2: 
            elements.append(Table([row2], colWidths=[13*cm]*len(row2), style=[('ALIGN', (0,0), (-1,-1), 'CENTER')]))
        
        doc.build(elements, onFirstPage=draw_header_footer, onLaterPages=draw_header_footer)
        app.after(0, lambda: ModernMessageBox.showinfo("Sucesso", "PDF gerado com sucesso!", parent=app))

    executar_com_carregamento(app, trabalho_pesado, "Gerando PDF", "Desenhando documento, aguarde...")

def gerar_solicitacao_nf_excel(app, record_ids, df_registros, banco_dict, template_path, callback_status, callback_atualizar):
    """Gera o Excel de solicitação de NF e atualiza o banco de dados."""
    df_selecionados = df_registros.loc[record_ids].copy()

    if len(df_selecionados['Parceiro'].unique()) > 1:
        ModernMessageBox.showerror("Erro", "Selecione registros de apenas UM Parceiro por vez.", parent=app)
        return

    def clean_str(x):
        s = str(x).strip() if pd.notna(x) else ""
        return "" if s.lower() in ("nan", "none") else s

    def fmt_moeda(val):
        return f"R$ {val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        
    nome_parceiro = df_selecionados['Parceiro'].iloc[0]
    
    dates = pd.to_datetime(df_selecionados['Data'], errors='coerce').dropna()
    if not dates.empty:
        data_carga_str = dates.max().strftime('%d/%m/%Y')
    else:
        data_carga_str = datetime.now().strftime('%d/%m/%Y')

    data_solicitacao_str = datetime.now().strftime("%d/%m/%Y")

    itens_agrupados = {}
    total_geral_carga = 0.0

    for _, row in df_selecionados.iterrows():
        nome_interno = clean_str(row.get('Residuo',''))
        dados_nf = banco_dict.get(nome_interno, {})
        desc_final = dados_nf.get('nome_nf', nome_interno)
        if not desc_final: desc_final = nome_interno
        cod_item = dados_nf.get('codigo_item', 'SEM_COD')
        
        valor_unit = float(row.get('ValorPorKg', 0))
        modo = str(row.get('Modo')).lower()
        
        if modo == 'fechado':
            chave = (cod_item, desc_final, "FECHADO") 
        else:
            chave = (cod_item, desc_final, valor_unit)

        if chave not in itens_agrupados:
            itens_agrupados[chave] = {'qtd': 0.0, 'total': 0.0, 'desc': desc_final, 'cod': cod_item, 'unit': valor_unit}
        
        peso_row = float(row.get('Peso', 0))
        val_row = float(row.get('ValorTotal', 0))
        
        itens_agrupados[chave]['qtd'] += peso_row
        itens_agrupados[chave]['total'] += val_row
        total_geral_carga += val_row

    if not os.path.exists(template_path):
        ModernMessageBox.showerror("Erro", f"Template '{template_path}' não encontrado.", parent=app)
        return

    try:
        wb = load_workbook(template_path)
        ws = wb.active
        
        # --- CORREÇÃO DO LOGO (#VALOR!) ---
        # 1. Apaga o erro das células principais no topo esquerdo
        ws['A1'].value = "" 
        
        try:
            # 2. Localiza o seu logo.png (mesma lógica do seu app principal)
            try: 
                base_path = sys._MEIPASS
            except Exception: 
                base_path = os.path.abspath(".")
                
            logo_path = os.path.join(base_path, "logo.png")
            
            # 3. Carimba a imagem na célula A1 do Excel
            if os.path.exists(logo_path):
                img_logo = XLImage(logo_path)
                img_logo.width = 310  # Ajuste a largura do logo no Excel aqui
                img_logo.height = 60  # Ajuste a altura do logo no Excel aqui
                ws.add_image(img_logo, 'A1')
        except Exception as e:
            print(f"Aviso: Não foi possível inserir o logo: {e}")
        # -----------------------------------
        
        placeholders = {
            '{{data_solicitacao}}': data_solicitacao_str,
            '{{nome_parceiro}}': nome_parceiro,
            '{{data_carga}}': data_carga_str,
            '{{valor_total_carga}}': fmt_moeda(total_geral_carga)
        }

        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    novo_texto = cell.value
                    alterou = False
                    for tag, valor_sub in placeholders.items():
                        if tag in novo_texto:
                            novo_texto = novo_texto.replace(tag, str(valor_sub))
                            alterou = True
                    if alterou:
                        cell.value = novo_texto

        start_row = -1
        col_map = {} 
        
        for r in range(1, 100):
            row_values = [str(c.value).strip() if c.value else "" for c in ws[r]]
            if "DESCRIÇÃO" in row_values and "VALOR" in row_values:
                start_row = r + 1
                for idx, txt in enumerate(row_values):
                    col_idx = idx + 1
                    if "Cód" in txt: col_map['cod'] = col_idx
                    if "QTDE" in txt or "PESO" in txt: col_map['qtde'] = col_idx
                    if "DESCRIÇÃO" in txt: col_map['desc'] = col_idx
                    if "VALOR" in txt or "UNIT" in txt: col_map['unit'] = col_idx
                    if "UN" in txt or "UN." in txt: col_map['un'] = col_idx
                break
        
        if start_row != -1 and col_map:
            current_row = start_row
            for chave, dados in itens_agrupados.items():
                if chave[2] == "FECHADO":
                    v_unit = 0
                    if dados['qtd'] > 0: v_unit = dados['total'] / dados['qtd']
                else:
                    v_unit = chave[2]

                if 'cod' in col_map: ws.cell(current_row, col_map['cod']).value = dados['cod']
                if 'un' in col_map: ws.cell(current_row, col_map['un']).value = "kg"
                
                if 'qtde' in col_map: 
                    c = ws.cell(current_row, col_map['qtde'])
                    c.value = dados['qtd']
                    c.number_format = '#,##0.000'
                
                if 'desc' in col_map: 
                    ws.cell(current_row, col_map['desc']).value = dados['desc']
                
                if 'unit' in col_map: 
                    c = ws.cell(current_row, col_map['unit'])
                    c.value = v_unit
                    c.number_format = '"R$" #,##0.00'

                current_row += 1

        data_arquivo = data_carga_str.replace("/", "-")
        import re as regex # Garantia de importação
        nome_limpo = str(nome_parceiro).strip()
        nome_limpo = regex.sub(r'[\\/*?:"<>|]', "", nome_limpo)
        filename = f"Solicitacao NF - {nome_limpo} - {data_arquivo}.xlsx"
        
        save_path = filedialog.asksaveasfilename(initialfile=filename, defaultextension=".xlsx", title="Salvar Solicitação de NF")
        
        if save_path:
            wb.save(save_path)
            
            if ModernMessageBox.askyesno("Sucesso", f"Arquivo gerado!\nTotal da Carga: {fmt_moeda(total_geral_carga)}\n\nDeseja marcar os itens como 'NF Solicitada'?", parent=app):
                for rid in record_ids:
                    try:
                        if rid in df_registros.index:
                            row = df_registros.loc[rid]
                            data_dt = row.get('Data', pd.NaT)
                            data_str = data_dt.strftime("%Y-%m-%d") if pd.notna(data_dt) else ""
                            data_full = [
                                data_str, clean_str(row.get('Parceiro', '')), clean_str(row.get('Residuo', '')),
                                clean_str(row.get('Estado', 'Sólido')), clean_str(row.get('Destinacao', '')),
                                clean_str(row.get('Tipo', 'Venda')), clean_str(row.get('Modo', 'kg')),
                                float(row.get('Peso', 0)), float(row.get('ValorPorKg', 0)),
                                float(row.get('ValorFechado', 0)), float(row.get('Transporte', 0)),
                                float(row.get('ValorTotal', 0)), clean_str(row.get('PedidoCompra', '')),
                                clean_str(row.get('NumMTR', '')), clean_str(row.get('CertificadoOK', 'Não')), 
                                float(row.get('PesoEmKg', 0)), "Solicitado", 
                                clean_str(row.get('CaminhoAnexo', ''))
                            ]
                            db.update_registro(rid, data_full)
                    except Exception as e: print(e)
                
                callback_atualizar()
                callback_status("Solicitação gerada e status atualizados.")
            else:
                callback_status("Solicitação gerada.")

    except Exception as e:
        ModernMessageBox.showerror("Erro Excel", f"{e}", parent=app)
